#!/usr/bin/env python3
"""
Build Product_on_Shelf_DB.xlsx — the backbone-database structure (DB_SCHEMA.md) as an
Excel file so it can be imported into the online Google Sheet without running Apps Script.

Mirrors Setup.gs exactly: same tabs, headers, and seed rows. Import into the existing
"Product on Shelf — Database" sheet via File -> Import -> Upload -> "Insert new sheet(s)".
Header row 1 is bold + frozen; the app reads by column NAME, so order is documentation only.
"""
from openpyxl import Workbook
from openpyxl.styles import Font

# Tab -> header row  (identical to DB_SCHEMA in Setup.gs)
SCHEMA = {
    "Spec": [
        "sku_key", "product", "vendor", "model", "category", "description",
        "spec_metric1_name", "spec_metric1_value",
        "spec_metric2_name", "spec_metric2_value",
        "spec_metric3_name", "spec_metric3_value",
        "unit", "citation_tag", "source", "source_detail", "date_time", "active",
    ],
    "Price": [
        "price_id", "sku_key", "product", "vendor", "model", "price_type",
        "amount_thb", "currency", "qty_basis", "effective_date", "superseded",
        "src_system", "src_ref", "citation_tag", "source", "source_detail", "date_time",
    ],
    "SOW": [
        "sow_id", "project", "impl_type", "product", "existing_equipment", "new_equipment",
        "sow_icr", "sla_tier", "support_years", "incidents", "site_prep",
        "impl_basic_thb", "impl_standard_thb", "impl_complex_thb",
        "citation_tag", "source", "source_detail", "date_time",
    ],
    "Estimates": [
        "estimate_id", "created_at", "created_by", "customer_name", "customer_detail",
        "request_price_thb", "buffer_applied", "status", "quotation_no", "notes",
    ],
    "Estimate_Lines": [
        "line_id", "estimate_id", "category", "vendor", "model", "sku_key",
        "qty", "unit_price_thb", "line_total_thb", "breakdown_json",
    ],
    "_meta": ["key", "value"],
    "_sync_log": [
        "run_id", "started_at", "finished_at", "src_system", "mode",
        "rows_scanned", "rows_upserted", "rows_superseded", "rows_skipped",
        "status", "error_detail", "chat_notified",
    ],
    "_alias": ["raw_string", "sku_key", "note"],
    "_catalog_map": ["product", "catalog_tab", "catalog_column", "mapping"],
}

TAB_ORDER = [
    "Spec", "Price", "SOW", "Estimates", "Estimate_Lines",
    "_meta", "_sync_log", "_alias", "_catalog_map",
]

META_SEED = [
    ["version", "1"],
    ["last_full_sync", ""],
    ["estimate_seq", "0"],
    ["email_backfill_done", "FALSE"],
    ["email_last_cursor", ""],
    ["owner", "oran.nin@scmtechnologies.co.th"],
]

# Firewall vertical slice — matches SHEET_SCHEMA.md Firewall tab columns
CATALOG_MAP_SEED = [
    ["firewall", "Firewall", "vendor", "field:vendor"],
    ["firewall", "Firewall", "model", "field:model"],
    ["firewall", "Firewall", "throughput_gbps", "spec:throughput_gbps"],
    ["firewall", "Firewall", "hw_thb", "price:hw"],
    ["firewall", "Firewall", "lic_yr_thb", "price:lic_yr"],
]

SEEDS = {"_meta": META_SEED, "_catalog_map": CATALOG_MAP_SEED}

wb = Workbook()
wb.remove(wb.active)  # drop default sheet

bold = Font(bold=True)
for name in TAB_ORDER:
    headers = SCHEMA[name]
    ws = wb.create_sheet(title=name)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        ws.cell(row=1, column=c).font = bold
    ws.freeze_panes = "A2"
    for row in SEEDS.get(name, []):
        ws.append(row)

out = "Product_on_Shelf_DB.xlsx"
wb.save(out)
print("wrote", out, "with tabs:", ", ".join(TAB_ORDER))
